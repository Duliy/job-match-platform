#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
match_api.py — 就业服务平台 统一服务 (FastAPI，校园交付版)

功能:
  - 用户注册/登录/认证
  - 岗位列表/统计/热门岗位 (公开，免登录)
  - AI 简历解析/岗位匹配/简历分析 (需登录 + 用户自带 DeepSeek API Key)
  - 岗位订阅推送 (SQL 粗筛，无需 AI Key)
  - 定时采集调度 + 管理后台
  - 前端静态页面服务 (单端口对外)

校园交付版说明:
  - 无 VIP/积分/付费体系，岗位浏览全部免费免登录
  - AI 功能使用学生自己填入的 DeepSeek Key（请求头 X-User-Api-Key 传递，
    仅存在学生浏览器 localStorage，服务器不留存）

启动:
    uvicorn match_api:app --host 0.0.0.0 --port 8000
"""

import hashlib
import json
import os
import time
from datetime import datetime
from typing import Optional, List
import io

from fastapi import FastAPI, File, UploadFile, Query, HTTPException, Depends, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
import PyPDF2
import pdfplumber
import fitz  # pymupdf — 最强PDF文本提取后备
import docx  # python-docx — Word简历提取
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

from match_db import (
    get_db,
    query_jobs,
    query_gk_jobs,
    get_all_cities,
    get_stats,
    pre_filter,
    record_hot_click,
)
from match_ai import (
    call_deepseek_match,
    parse_resume_with_deepseek,
    analyze_resume_with_deepseek,
)
from auth import (
    register_user,
    authenticate_user,
    get_current_user,
    get_optional_user,
    log_action,
    get_email_config,
    save_email_config,
)
from admin_api import router as admin_router


def _extract_pdf_text(content: bytes) -> str:
    """
    从PDF字节内容中提取文本，依次尝试三种方案：
    1. PyPDF2（速度快，对英文/标准PDF好）
    2. pdfplumber（对中文/复杂布局好）
    3. pymupdf/fitz（最强后备，支持最广）
    全部失败则返回空字符串（说明是扫描件/纯图片PDF，需OCR）
    """
    # 方案1: PyPDF2（速度快）
    raw_text = ""
    try:
        pdf_file = io.BytesIO(content)
        reader = PyPDF2.PdfReader(pdf_file)
        for page in reader.pages:
            text = page.extract_text()
            if text:
                raw_text += text + "\n"
    except Exception:
        pass

    if raw_text.strip():
        return raw_text.strip()

    # 方案2: pdfplumber（对中文PDF支持好）
    try:
        pdf_file2 = io.BytesIO(content)
        with pdfplumber.open(pdf_file2) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    raw_text += text + "\n"
    except Exception:
        pass

    if raw_text.strip():
        return raw_text.strip()

    # 方案3: pymupdf/fitz（最强后备）
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        for page in doc:
            text = page.get_text()
            if text:
                raw_text += text + "\n"
        doc.close()
    except Exception:
        pass

    return raw_text.strip()


def _extract_docx_text(content: bytes) -> str:
    """
    从Word (.docx) 字节内容中提取文本。
    先用 python-docx 的 paragraph/table 提取，如果为空再用
    lxml 直接从 XML 中提取所有 w:t 文本（兼容特殊嵌套结构）。
    python-docx 仅支持 .docx (Office 2007+) 格式。
    失败时返回空字符串。
    """
    try:
        docx_file = io.BytesIO(content)
        doc = docx.Document(docx_file)
        paragraphs = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                paragraphs.append(t)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        paragraphs.append(t)

        result = "\n".join(paragraphs)
        if result.strip():
            return result

        # python-docx 提取为空 → 直接解析 XML 提取所有 w:t 文本
        # 适用于使用了特殊嵌套/SDT/ContentControl 的简历模板
        from lxml import etree

        ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        t_elements = doc.element.findall(".//{%s}t" % ns_w)
        raw_texts = []
        for t_elem in t_elements:
            if t_elem.text:
                raw_texts.append(t_elem.text)
        return "\n".join(raw_texts)
    except Exception:
        return ""


def _extract_resume_text(content: bytes, filename: str) -> str:
    """
    根据文件类型智能提取简历文本。
    支持 .pdf（三重PDF引擎）和 .docx（Word文档）。
    .doc 格式需要先用 WPS/Word 另存为 .docx。
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".doc":
        # 旧版 .doc (OLE) 格式，python-docx 不支持
        raise HTTPException(
            400,
            "旧版.doc格式不支持，请用WPS/Word打开后「另存为」→ 选择「Word文档(.docx)」格式再上传",
        )
    elif ext == ".docx":
        return _extract_docx_text(content)
    else:
        return _extract_pdf_text(content)


# ── 配置 ──────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "jobs.db")
CACHE_TTL = 300  # 缓存5分钟

# 对外访问地址（用于邮件中的链接）。部署时由启动脚本自动写入，
# 例如 http://192.168.1.100:8000 ；未配置时邮件中不放链接。
PUBLIC_BASE_URL = os.environ.get("PUBLIC_BASE_URL", "").rstrip("/")

# ── FastAPI 应用 ─────────────────────────────────────
app = FastAPI(title="就业服务平台API（校园版）", version="3.0.0", docs_url="/api/docs")

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 挂载管理后台路由
app.include_router(admin_router)


# ══════════════════════════════════════════════════════
# 用户自带 API Key 解析（校园版核心机制）
# ══════════════════════════════════════════════════════
def resolve_user_api_key(request: Request) -> str:
    """从请求头 X-User-Api-Key 获取学生自己的 DeepSeek Key。

    Key 由学生保存在自己浏览器的 localStorage 中，随请求带来，
    服务器只用于本次调用，绝不落库存储。
    找不到时回退到服务器环境变量 DEEPSEEK_API_KEY（管理员可选配置）。
    """
    key = (request.headers.get("x-user-api-key") or "").strip()
    if key:
        return key
    env_key = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if env_key:
        return env_key
    raise HTTPException(
        status_code=402,
        detail="NO_API_KEY: 使用 AI 功能需要先填入你自己的 DeepSeek API Key"
        "（页面顶部「API Key 设置」中填写，Key 仅保存在你自己的浏览器里）",
    )


# 前端页面与静态资源（项目根目录下的白名单文件）
_STATIC_FILES = {
    "index.html": "text/html; charset=utf-8",
    "resume_match.html": "text/html; charset=utf-8",
    "admin.html": "text/html; charset=utf-8",
    "manifest.json": "application/json; charset=utf-8",
    "sw.js": "application/javascript; charset=utf-8",
    "icon-192.png": "image/png",
    "icon-512.png": "image/png",
}


# ══════════════════════════════════════════════════════
# 根路由 — 前端入口页（学生直接访问 http://服务器IP/ 即可）
# ══════════════════════════════════════════════════════
@app.get("/", include_in_schema=False)
async def root():
    return FileResponse(
        os.path.join(BASE_DIR, "index.html"),
        headers={"Cache-Control": "no-cache"},
    )


# ══════════════════════════════════════════════════════
# Pydantic 数据模型
# ══════════════════════════════════════════════════════
class RegisterRequest(BaseModel):
    username: str
    password: str


class LoginRequest(BaseModel):
    username: str
    password: str


class ResumeInfo(BaseModel):
    """简历解析结果"""

    name: str = ""
    education: str = ""
    major: str = ""
    school: str = ""
    skills: List[str] = []
    experience: str = ""
    raw_text: str = ""


class MatchPreference(BaseModel):
    """用户匹配偏好"""

    cities: List[str] = []
    salary_min: float = 0
    salary_max: float = 999999
    job_type: str = "社招"
    education_match: bool = True


class MatchRequest(BaseModel):
    """匹配请求"""

    resume: ResumeInfo
    preference: MatchPreference


class JobRecommendation(BaseModel):
    """单个岗位推荐结果"""

    job_id: int = 0
    job_type: str = "社招"
    title: str = ""
    company: str = ""
    salary: str = ""
    salary_raw: str = ""
    salary_min: float = 0
    salary_max: float = 0
    location: str = ""
    education: str = ""
    platform: str = ""
    link: str = ""
    match_score: int = 0
    admission_probability: str = "中"
    reason: str = ""
    advantages: List[str] = []
    suggestions: List[str] = []


class MatchResponse(BaseModel):
    """匹配响应"""

    recommendations: List[JobRecommendation] = []
    total_candidates: int = 0
    ai_analyzed_count: int = 0
    cached: bool = False
    analysis_time_ms: float = 0


# ══════════════════════════════════════════════════════
# 认证 API
# ══════════════════════════════════════════════════════


@app.post("/api/auth/register")
async def register(req: RegisterRequest):
    """用户注册"""
    result = register_user(req.username, req.password)
    if not result["ok"]:
        raise HTTPException(400, result["message"])

    # 自动登录
    login_result = authenticate_user(req.username, req.password)
    return {
        "ok": True,
        "user_id": result["user_id"],
        "token": login_result.get("token", ""),
        "username": req.username,
    }


@app.post("/api/auth/login")
async def login(req: LoginRequest):
    """用户登录"""
    result = authenticate_user(req.username, req.password)
    if not result["ok"]:
        raise HTTPException(401, result["message"])

    return {
        "ok": True,
        "token": result["token"],
        "user_id": result["user_id"],
        "username": req.username,
        "role": result["role"],
    }


@app.get("/api/auth/me")
async def get_me(user: dict = Depends(get_current_user)):
    """获取当前用户信息"""
    return {
        "id": user["id"],
        "username": user["username"],
        "role": user["role"],
        "created_at": user.get("created_at", ""),
    }


@app.get("/api/auth/verify")
async def verify_token(user: dict = Depends(get_current_user)):
    """校验 token 有效性（管理后台自动登录用）"""
    return {"ok": True, "username": user["username"], "role": user["role"]}


class ChangePasswordRequest(BaseModel):
    old_password: str
    new_password: str


@app.post("/api/auth/change-password")
async def change_password(
    req: ChangePasswordRequest, user: dict = Depends(get_current_user)
):
    """修改密码（管理后台「系统设置」用）"""
    if len(req.new_password) < 6:
        raise HTTPException(400, "新密码至少6位")
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT password_hash FROM users WHERE id = ?", (user["id"],))
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "用户不存在")
    from auth import verify_password, hash_password

    if not verify_password(req.old_password, row["password_hash"]):
        conn.close()
        raise HTTPException(400, "当前密码错误")
    cursor.execute(
        "UPDATE users SET password_hash = ? WHERE id = ?",
        (hash_password(req.new_password), user["id"]),
    )
    conn.commit()
    conn.close()
    log_action(user["id"], "change_password", "")
    return {"ok": True, "message": "密码已修改"}


# ══════════════════════════════════════════════════════
# 热力图 API
# ══════════════════════════════════════════════════════


@app.get("/api/hot/jobs")
async def hot_jobs(limit: int = Query(20, ge=1, le=100)):
    """获取热门岗位列表（公开，无需登录）"""
    from match_db import get_hot_jobs

    return get_hot_jobs(limit)


@app.post("/api/hot/click/{job_id}")
async def hot_click(
    job_id: int, job_type: str = Query("社招"), user: dict = Depends(get_optional_user)
):
    """记录热门岗位点击"""
    record_hot_click(job_id, job_type)
    if user:
        log_action(
            user["id"],
            "click_job",
            json.dumps({"job_id": job_id, "job_type": job_type}),
        )
    return {"ok": True}


# ══════════════════════════════════════════════════════
# AI 功能 — 简历解析（需登录 + 用户自带 DeepSeek Key）
# ══════════════════════════════════════════════════════


@app.post("/api/match/parse-resume", response_model=ResumeInfo)
async def parse_resume(
    request: Request,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    """
    接收PDF或Word简历文件，返回解析后的简历结构。
    校园版：使用请求头 X-User-Api-Key 中学生自己的 DeepSeek Key，服务器不扣费不留存。
    """
    api_key = resolve_user_api_key(request)

    # 校验文件类型
    fname = file.filename.lower()
    if not (
        fname.endswith(".pdf") or fname.endswith(".docx") or fname.endswith(".doc")
    ):
        raise HTTPException(400, "仅支持PDF或Word(.docx/.doc)格式简历")

    # 读取文件内容
    content = await file.read()

    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    # 提取文本
    raw_text = _extract_resume_text(content, file.filename)

    if not raw_text.strip():
        fname = file.filename.lower()
        if fname.endswith(".pdf"):
            hint = "（可能原因：1.PDF为扫描件/图片格式；2.PDF加密或损坏。建议用WPS打开后「另存为」可搜索PDF，或改用Word(.docx)格式上传）"
        else:
            hint = "（可能原因：1.Word文档为空或只含图片；2.文件损坏。建议检查文件内容后重新上传）"
        raise HTTPException(400, "无法从文件中提取文本内容" + hint)

    # 调用 DeepSeek 解析（使用学生自己的 Key）
    try:
        resume_info, usage = await parse_resume_with_deepseek(raw_text, api_key=api_key)
        resume_info["raw_text"] = raw_text

        # 记录行为
        log_action(user["id"], "resume_parse", json.dumps({"filename": file.filename}))

        return ResumeInfo(**resume_info)
    except ValueError as e:
        raise HTTPException(500, f"AI服务配置错误: {str(e)}")
    except Exception as e:
        raise HTTPException(500, f"简历解析失败: {str(e)}")


# ══════════════════════════════════════════════════════
# AI 功能 — 岗位匹配（需登录 + 用户自带 DeepSeek Key）
# ══════════════════════════════════════════════════════


@app.post("/api/match/recommend", response_model=MatchResponse)
async def recommend(
    req: MatchRequest,
    request: Request,
    user: dict = Depends(get_current_user),
):
    """
    简历 + 偏好 → 岗位推荐。校园版：使用学生自己的 DeepSeek Key（请求头 X-User-Api-Key）。
    """
    api_key = resolve_user_api_key(request)

    start_time = time.time()

    # 1. 缓存检查
    resume_hash = hashlib.md5(req.resume.raw_text.encode()).hexdigest()
    pref_hash = hashlib.md5(
        json.dumps(req.preference.model_dump(), sort_keys=True).encode()
    ).hexdigest()

    cached = check_cache(resume_hash, pref_hash, req.preference.job_type)
    if cached:
        return MatchResponse(**cached, cached=True, analysis_time_ms=0)

    # 2. 预筛选
    candidates = pre_filter(
        skills=req.resume.skills,
        cities=req.preference.cities,
        salary_min=req.preference.salary_min,
        salary_max=req.preference.salary_max,
        education=req.resume.education,
        job_type=req.preference.job_type,
        max_candidates=100,
    )

    if not candidates:
        candidates = pre_filter(
            skills=[],
            cities=req.preference.cities,
            salary_min=req.preference.salary_min,
            salary_max=req.preference.salary_max,
            education="",
            job_type=req.preference.job_type,
            max_candidates=50,
        )

    if not candidates:
        return MatchResponse(
            recommendations=[],
            total_candidates=0,
            ai_analyzed_count=0,
            cached=False,
            analysis_time_ms=(time.time() - start_time) * 1000,
        )

    # 3. DeepSeek 分析（使用学生自己的 Key）
    resume_dict = req.resume.model_dump()
    pref_dict = req.preference.model_dump()

    try:
        ai_results, usage = await call_deepseek_match(
            resume_dict, pref_dict, candidates, api_key=api_key
        )
    except ValueError as e:
        raise HTTPException(500, f"AI服务配置错误: {str(e)}")
    except Exception:
        ai_results = []
        for job in candidates[:10]:
            ai_results.append(
                {
                    "job_id": job.get("id", 0),
                    "match_score": 50,
                    "admission_probability": "中",
                    "reason": "AI分析暂时不可用，此为默认排序",
                    "advantages": [],
                    "suggestions": [],
                }
            )
        usage = {"input_tokens": 0, "output_tokens": 0}

    # 4. 合并AI结果与数据库岗位信息
    job_map = {job["id"]: job for job in candidates}
    recommendations = []
    for ai_rec in ai_results:
        job_id = ai_rec.get("job_id", 0)
        job = job_map.get(job_id, {})
        rec = JobRecommendation(
            job_id=job_id,
            job_type=job.get("job_type", req.preference.job_type),
            title=job.get("title", ""),
            company=job.get("company", ""),
            salary=job.get("salary", ""),
            salary_raw=job.get("salary_raw", "面议"),
            salary_min=job.get("salary_min", 0),
            salary_max=job.get("salary_max", 0),
            location=job.get("location", "") or job.get("province", ""),
            education=job.get("education", ""),
            platform=job.get("platform", "") or job.get("source", ""),
            link=job.get("link", "") or job.get("position_id", ""),
            match_score=ai_rec.get("match_score", 0),
            admission_probability=ai_rec.get("admission_probability", "中"),
            reason=ai_rec.get("reason", ""),
            advantages=ai_rec.get("advantages", []),
            suggestions=ai_rec.get("suggestions", []),
        )
        recommendations.append(rec)

    # 5. 写入缓存
    result_data = {
        "recommendations": [r.model_dump() for r in recommendations],
        "total_candidates": len(candidates),
        "ai_analyzed_count": len(recommendations),
    }
    save_cache(resume_hash, pref_hash, req.preference.job_type, result_data)

    # 记录行为
    log_action(
        user["id"],
        "ai_match",
        json.dumps({"candidates": len(candidates), "results": len(recommendations)}),
    )

    elapsed_ms = (time.time() - start_time) * 1000
    return MatchResponse(
        recommendations=recommendations,
        total_candidates=len(candidates),
        ai_analyzed_count=len(recommendations),
        cached=False,
        analysis_time_ms=round(elapsed_ms, 0),
    )


# ══════════════════════════════════════════════════════
# AI 功能 — 简历分析（需登录 + 用户自带 DeepSeek Key）
# ══════════════════════════════════════════════════════


@app.post("/api/match/analyze-resume")
async def analyze_resume(
    request: Request,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    """
    上传PDF或Word简历，AI深度分析点评。校园版：使用学生自己的 DeepSeek Key。
    """
    api_key = resolve_user_api_key(request)

    # 校验文件类型
    fname = file.filename.lower()
    if not (
        fname.endswith(".pdf") or fname.endswith(".docx") or fname.endswith(".doc")
    ):
        raise HTTPException(400, "仅支持PDF或Word(.docx/.doc)格式简历")

    # 读取文件内容
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    # 提取文本
    raw_text = _extract_resume_text(content, file.filename)

    if not raw_text.strip():
        fname = file.filename.lower()
        if fname.endswith(".pdf"):
            hint = "（可能原因：1.PDF为扫描件/图片格式；2.PDF加密或损坏。建议用WPS打开后「另存为」可搜索PDF，或改用Word(.docx)格式上传）"
        else:
            hint = "（可能原因：1.Word文档为空或只含图片；2.文件损坏。建议检查文件内容后重新上传）"
        raise HTTPException(400, "无法从文件中提取文本内容" + hint)

    # 调用 DeepSeek 分析（使用学生自己的 Key）
    try:
        analysis_result, usage = await analyze_resume_with_deepseek(
            raw_text, api_key=api_key
        )

        # 记录行为
        log_action(user["id"], "ai_analysis", json.dumps({"filename": file.filename}))

        return {"ok": True, "analysis": analysis_result}
    except ValueError as e:
        raise HTTPException(500, f"AI服务配置错误: {str(e)}")
    except Exception as e:
        raise HTTPException(500, f"简历分析失败: {str(e)}")


# ══════════════════════════════════════════════════════
# 公开 API（无需登录）
# ══════════════════════════════════════════════════════


@app.get("/api/match/jobs")
async def list_jobs(
    keyword: str = Query(""),
    city: str = Query(""),
    salary_min: float = Query(0),
    salary_max: float = Query(999999),
    job_type: str = Query("社招"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    """岗位列表 (公开，分页 + 筛选)"""
    if job_type == "公考":
        rows, total = query_gk_jobs(keyword, city, page, page_size)
    else:
        rows, total = query_jobs(keyword, city, salary_min, salary_max, page, page_size)

    return {
        "items": rows,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": (total + page_size - 1) // page_size if total > 0 else 0,
    }


@app.get("/api/match/cities")
async def list_cities(job_type: str = Query("社招")):
    """全库城市/省份列表 (公开，供筛选下拉)"""
    return {"cities": get_all_cities(job_type)}


@app.get("/api/match/stats")
async def stats():
    """数据库统计信息 (公开)"""
    return get_stats()


# ══════════════════════════════════════════════════════
# 订阅与简历托管 API（校园版：全部登录用户可用，无会员门槛）
# 说明：路径保留 /api/vip/* 以兼容已有前端，语义已变为"订阅服务"。
# ══════════════════════════════════════════════════════


@app.post("/api/vip/resume/upload")
async def vip_upload_resume(
    request: Request,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    """上传/更新简历（用于订阅推送的粗筛依据）。

    校园版：所有登录用户可用。若请求头带有学生自己的 DeepSeek Key，
    顺便做一次 AI 结构化解析（技能/学历），用于订阅推送的 SQL 粗筛；
    没有 Key 也能上传，只是推送匹配精度会粗一些。
    """

    # 校验文件类型
    fname = file.filename.lower()
    if not (
        fname.endswith(".pdf") or fname.endswith(".docx") or fname.endswith(".doc")
    ):
        raise HTTPException(400, "仅支持PDF或Word(.docx/.doc)格式简历")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    raw_text = _extract_resume_text(content, file.filename)
    if not raw_text.strip():
        raise HTTPException(400, "无法从文件中提取文本内容，请检查文件格式")

    # 若学生带了自己的 Key，顺便做 AI 结构化解析（技能/学历），用于订阅粗筛
    parsed_json = "{}"
    user_key = (request.headers.get("x-user-api-key") or "").strip()
    if user_key:
        try:
            parsed_info, _usage = await parse_resume_with_deepseek(
                raw_text, api_key=user_key
            )
            parsed_info.pop("raw_text", None)
            parsed_json = json.dumps(parsed_info, ensure_ascii=False)
        except Exception:
            pass  # 解析失败不阻塞上传，粗筛会退化为宽泛匹配

    # 存入 member_resumes
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM member_resumes WHERE user_id = ?", (user["id"],))
    existing = cursor.fetchone()

    now = datetime.now().isoformat()
    if existing:
        cursor.execute(
            """
            UPDATE member_resumes
            SET resume_filename=?, resume_text=?, parsed_json=?, analysis_json='{}', updated_at=?
            WHERE user_id=?
        """,
            (file.filename, raw_text, parsed_json, now, user["id"]),
        )
        resume_id = existing["id"]
    else:
        cursor.execute(
            """
            INSERT INTO member_resumes (user_id, resume_filename, resume_text, parsed_json, updated_at)
            VALUES (?, ?, ?, ?, ?)
        """,
            (user["id"], file.filename, raw_text, parsed_json, now),
        )
        resume_id = cursor.lastrowid

    conn.commit()
    conn.close()

    log_action(user["id"], "vip_upload_resume", json.dumps({"filename": file.filename}))
    return {
        "ok": True,
        "resume_id": resume_id,
        "text_length": len(raw_text),
        "message": "简历上传成功",
    }


@app.get("/api/vip/resume")
async def vip_get_resume(user: dict = Depends(get_current_user)):
    """获取已存储的简历（校园版：所有登录用户可用）"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, resume_filename, resume_text, parsed_json, analysis_json, uploaded_at, updated_at FROM member_resumes WHERE user_id = ?",
        (user["id"],),
    )
    row = cursor.fetchone()
    conn.close()

    if not row:
        return {"ok": True, "has_resume": False, "message": "尚未上传简历"}

    return {
        "ok": True,
        "has_resume": True,
        "resume_id": row["id"],
        "filename": row["resume_filename"],
        "text_length": len(row["resume_text"]),
        "parsed": json.loads(row["parsed_json"]) if row["parsed_json"] else {},
        "analysis": json.loads(row["analysis_json"]) if row["analysis_json"] else {},
        "uploaded_at": row["uploaded_at"],
        "updated_at": row["updated_at"],
    }


@app.post("/api/vip/resume/analyze")
async def vip_analyze_resume(request: Request, user: dict = Depends(get_current_user)):
    """对已上传简历做 AI 深度分析（使用学生自己的 DeepSeek Key）"""
    api_key = resolve_user_api_key(request)

    # 获取存储的简历文本
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, resume_text, analysis_json FROM member_resumes WHERE user_id = ?",
        (user["id"],),
    )
    row = cursor.fetchone()
    if not row or not row["resume_text"].strip():
        conn.close()
        raise HTTPException(400, "请先上传简历")

    resume_id = row["id"]
    resume_text = row["resume_text"]

    # 检查是否已有分析结果（7天内有效）
    existing = json.loads(row["analysis_json"]) if row["analysis_json"] else {}
    if existing and existing.get("analyzed_at"):
        try:
            analyzed_dt = datetime.fromisoformat(existing["analyzed_at"])
            if (datetime.now() - analyzed_dt).days < 7:
                conn.close()
                return {"ok": True, "cached": True, "analysis": existing}
        except Exception:
            pass

    conn.close()

    # 调用 DeepSeek 分析（使用学生自己的 Key）
    try:
        analysis_result, usage = await analyze_resume_with_deepseek(
            resume_text, api_key=api_key
        )
        analysis_result["analyzed_at"] = datetime.now().isoformat()
    except ValueError as e:
        raise HTTPException(500, f"AI服务配置错误: {str(e)}")
    except Exception as e:
        raise HTTPException(500, f"简历分析失败: {str(e)}")

    # 保存分析结果
    conn_db = get_db()
    conn_db.cursor().execute(
        "UPDATE member_resumes SET analysis_json=? WHERE id=?",
        (json.dumps(analysis_result, ensure_ascii=False), resume_id),
    )
    conn_db.commit()
    conn_db.close()

    log_action(user["id"], "vip_analyze_resume", json.dumps({"resume_id": resume_id}))
    return {"ok": True, "cached": False, "analysis": analysis_result}


@app.post("/api/vip/match")
async def vip_match_jobs(request: Request, user: dict = Depends(get_current_user)):
    """对已上传简历做岗位 AI 匹配（使用学生自己的 DeepSeek Key）"""
    api_key = resolve_user_api_key(request)

    # 获取存储的简历
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, resume_text, parsed_json FROM member_resumes WHERE user_id = ?",
        (user["id"],),
    )
    row = cursor.fetchone()
    if not row or not row["resume_text"].strip():
        conn.close()
        raise HTTPException(400, "请先上传简历")

    resume_id = row["id"]
    resume_text = row["resume_text"]
    parsed = json.loads(row["parsed_json"]) if row["parsed_json"] else {}

    # 尝试用解析结果
    if parsed and parsed.get("skills"):
        skills = parsed.get("skills", [])
        education = parsed.get("education", "")
    else:
        # 先快速解析（用学生自己的 Key）
        try:
            resume_info, _use = await parse_resume_with_deepseek(
                resume_text, api_key=api_key
            )
            skills = resume_info.get("skills", [])
            education = resume_info.get("education", "")
            parsed = resume_info
            parsed["raw_text"] = resume_text
            # 保存解析结果
            cursor.execute(
                "UPDATE member_resumes SET parsed_json=? WHERE id=?",
                (json.dumps(resume_info, ensure_ascii=False), resume_id),
            )
            conn.commit()
        except Exception:
            skills = []
            education = ""

    conn.close()

    start_time = time.time()

    # 累计 token 用量（用于成本统计）
    total_input_tokens = 0
    total_output_tokens = 0

    # 预筛选 (社招+公考各搜一次) — 收集所有 candidates 用于合并 AI 结果
    all_recommendations = []
    all_candidates_map = {}  # job_id → 原始岗位数据（用于补充 title/company/salary 等）

    for jt in ["社招", "公考"]:
        candidates = pre_filter(
            skills=skills,
            cities=[],
            salary_min=0,
            salary_max=999999,
            education=education,
            job_type=jt,
            max_candidates=50,
        )
        # 技能词是英文（如Python/Java），数据库关键字是中文（如计算机/公务员），可能匹配不上
        # 回退到宽泛搜索，把岗位拉出来让AI打分排序
        if not candidates and skills:
            candidates = pre_filter(
                skills=[],
                cities=[],
                salary_min=0,
                salary_max=999999,
                education=education,
                job_type=jt,
                max_candidates=50,
            )
        if not candidates:
            continue

        # 建立 job_id → 原始数据映射（AI 结果不含 title/company/salary/link 等字段）
        for c in candidates:
            cid = c.get("id")
            if cid and cid not in all_candidates_map:
                all_candidates_map[cid] = c

        resume_dict = {
            "skills": skills,
            "education": education,
            "experience": "",
            "name": "",
            "major": "",
            "school": "",
            "raw_text": resume_text,
        }
        pref_dict = {
            "cities": [],
            "salary_min": 0,
            "salary_max": 999999,
            "job_type": jt,
            "education_match": True,
        }

        try:
            ai_results, usage = await call_deepseek_match(
                resume_dict, pref_dict, candidates, api_key=api_key
            )
            total_input_tokens += usage.get("input_tokens", 0)
            total_output_tokens += usage.get("output_tokens", 0)
        except Exception:
            ai_results = []
            for c in candidates[:10]:
                ai_results.append(
                    {
                        "job_id": c.get("id", 0),
                        "job_type": jt,
                        "match_score": 50,
                        "admission_probability": "中",
                        "reason": "AI服务暂时不可用，显示默认推荐",
                        "advantages": [],
                        "suggestions": [],
                    }
                )

        all_recommendations.extend(ai_results)

    all_recommendations.sort(key=lambda x: x.get("match_score", 0), reverse=True)
    top15 = all_recommendations[:15]

    # 格式化结果 — 从原始 candidates 补充缺失的岗位详情字段
    recommendations = []
    high_prob_count = 0
    for r in top15:
        jid = r.get("job_id", 0)
        src = all_candidates_map.get(jid, {})
        rec = JobRecommendation(
            job_id=jid,
            job_type=r.get("job_type", "社招") or src.get("job_type", "社招"),
            title=r.get("title", "") or src.get("title", ""),
            company=r.get("company", "") or src.get("company", ""),
            salary=r.get("salary", "")
            or src.get("salary", "")
            or src.get("salary_display", ""),
            salary_raw=r.get("salary_raw", "") or src.get("salary_raw", ""),
            salary_min=r.get("salary_min", 0) or src.get("salary_min", 0),
            salary_max=r.get("salary_max", 0) or src.get("salary_max", 0),
            location=r.get("location", "")
            or src.get("city", "")
            or src.get("location", ""),
            education=r.get("education", "") or src.get("education", ""),
            platform=r.get("platform", "") or src.get("platform", ""),
            link=r.get("link", "") or src.get("link", "") or src.get("url", ""),
            match_score=r.get("match_score", 0),
            admission_probability=r.get("admission_probability", "中"),
            reason=r.get("reason", ""),
            advantages=r.get("advantages", []),
            suggestions=r.get("suggestions", []),
        )
        recommendations.append(rec)
        if r.get("match_score", 0) >= 60:
            high_prob_count += 1

    # 保存匹配结果
    result_json = json.dumps(
        [r.model_dump() for r in recommendations], ensure_ascii=False
    )
    match_db = get_db()
    match_db.cursor().execute(
        "INSERT INTO vip_match_results (user_id, resume_id, match_json, high_prob_count) VALUES (?, ?, ?, ?)",
        (user["id"], resume_id, result_json, high_prob_count),
    )
    match_db.commit()
    match_db.close()

    # 如果有高匹配度岗位(>=60%)且用户绑定了通知邮箱，立即发送邮件
    if high_prob_count > 0:
        _send_match_email_async(user, high_prob_count, recommendations)

    analysis_time = (time.time() - start_time) * 1000
    log_action(
        user["id"],
        "vip_match",
        json.dumps(
            {
                "resume_id": resume_id,
                "count": len(recommendations),
                "high_prob": high_prob_count,
            }
        ),
    )

    return {
        "ok": True,
        "recommendations": [r.model_dump() for r in recommendations],
        "total_candidates": len(top15),
        "high_prob_count": high_prob_count,
        "analysis_time_ms": analysis_time,
    }


@app.get("/api/vip/match/results")
async def vip_match_results(user: dict = Depends(get_current_user)):
    """获取最近的匹配结果（校园版：所有登录用户可用）"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT id, match_json, high_prob_count, notified, created_at
        FROM vip_match_results WHERE user_id = ? ORDER BY created_at DESC LIMIT 5
    """,
        (user["id"],),
    )
    rows = cursor.fetchall()
    conn.close()

    results = []
    for r in rows:
        try:
            matches = json.loads(r["match_json"])
        except Exception:
            matches = []
        results.append(
            {
                "id": r["id"],
                "matches": matches,
                "high_prob_count": r["high_prob_count"],
                "notified": bool(r["notified"]),
                "created_at": r["created_at"],
            }
        )

    return {"ok": True, "results": results}


# ══════════════════════════════════════════════════════
# 邮件配置 API (管理员)
# ══════════════════════════════════════════════════════


class EmailConfigRequest(BaseModel):
    smtp_host: str = ""
    smtp_port: int = 465
    smtp_user: str = ""
    smtp_pass: str = ""
    from_email: str = ""
    enabled: int = 0


@app.get("/api/vip/email-config")
async def get_email_cfg(user: dict = Depends(get_current_user)):
    """获取邮件配置（管理员）"""
    if user.get("role") != "admin":
        return {"ok": True, "enabled": 0, "message": "仅管理员可查看完整配置"}
    return {"ok": True, "config": get_email_config()}


@app.post("/api/vip/email-config")
async def save_email_cfg(
    req: EmailConfigRequest, user: dict = Depends(get_current_user)
):
    """更新邮件配置（管理员）"""
    if user.get("role") != "admin":
        raise HTTPException(403, "需要管理员权限")
    result = save_email_config(req.model_dump())
    return result


class NotifyEmailRequest(BaseModel):
    email: str


@app.post("/api/vip/notify-email")
async def set_notify_email(
    req: NotifyEmailRequest, user: dict = Depends(get_current_user)
):
    """Set the user's notification email address"""
    if "@" not in req.email or "." not in req.email:
        raise HTTPException(400, "Invalid email format")
    conn = get_db()
    conn.execute("UPDATE users SET notify_email=? WHERE id=?", (req.email, user["id"]))
    conn.commit()
    conn.close()
    return {"ok": True, "message": "Notification email updated"}


@app.get("/api/vip/notify-email")
async def get_notify_email(user: dict = Depends(get_current_user)):
    """Get the user's notification email address"""
    conn = get_db()
    row = conn.execute(
        "SELECT notify_email FROM users WHERE id=?", (user["id"],)
    ).fetchone()
    conn.close()
    return {"ok": True, "email": row["notify_email"] if row else ""}


def _send_match_email_async(user: dict, high_prob_count: int, recommendations):
    """异步发送匹配结果邮件通知（不阻塞响应）"""
    import threading
    import os

    log_path = os.path.join(os.path.dirname(__file__), "data", "email_debug.log")

    def _log(msg):
        from datetime import datetime as dt

        line = f"[{dt.now().strftime('%H:%M:%S')}] {msg}\n"
        print(f"[MatchEmail] {msg}")
        try:
            with open(log_path, "a", encoding="utf-8") as f:
                f.write(line)
        except Exception:
            pass

    def _do_send():
        try:
            _log(
                f"Starting email send for user={user.get('username')} high_prob={high_prob_count} recs_count={len(recommendations)}"
            )

            # 1. 获取用户通知邮箱
            conn = get_db()
            row = conn.execute(
                "SELECT notify_email FROM users WHERE id=?", (user["id"],)
            ).fetchone()
            notify_email = (row["notify_email"] or "") if row else ""
            if not notify_email:
                _log(f"User {user['username']} has no notify_email set, skip")
                conn.close()
                return
            _log(f"notify_email={notify_email}")

            # 2. 获取SMTP配置
            email_cfg = get_email_config()
            if not email_cfg or not email_cfg.get("enabled"):
                _log(f"SMTP not enabled (cfg={email_cfg}), skip")
                conn.close()
                return
            _log(
                f"SMTP configured: {email_cfg['smtp_host']}:{email_cfg['smtp_port']}, from={email_cfg['from_email']}"
            )
            conn.close()

            # 3. 筛选高匹配度岗位(>=60%) — 兼容 dict 和 Pydantic 模型
            high_jobs = []
            for r in recommendations:
                score = getattr(r, "match_score", None)
                if score is None:
                    score = r.get("match_score", 0) if isinstance(r, dict) else 0
                if score >= 60:
                    high_jobs.append(r)

            _log(f"Filtered high_jobs: {len(high_jobs)} jobs (threshold=60)")
            if not high_jobs:
                _log("No jobs above threshold, skip sending")
                return

            # 4. 构建HTML卡片邮件 — 参考 test_smtp.py 的简洁渐变卡片风格
            def _jattr(obj, key, default=""):
                v = getattr(obj, key, None)
                if v is not None:
                    return str(v)
                if isinstance(obj, dict):
                    return str(obj.get(key, default))
                return default

            now_str = datetime.now().strftime("%Y年%m月d日 %H:%M")
            job_count = len(high_jobs)

            # 生成岗位列表
            jobs_html = ""
            for j in high_jobs[:10]:
                score = int(_jattr(j, "match_score", 0))
                title = _jattr(j, "title", "未知职位")
                company = _jattr(j, "company", "")
                salary = _jattr(j, "salary", "") or _jattr(j, "salary_raw", "")
                location = _jattr(j, "location", "")
                platform = _jattr(j, "platform", "")
                link = _jattr(j, "link", "")
                reason = _jattr(j, "reason", "")

                if score >= 80:
                    score_bg = "#E8F5EE"
                    score_color = "#1B7343"
                else:
                    score_bg = "#FBF3E5"
                    score_color = "#9C6B3C"

                detail_lines = []
                if company:
                    detail_lines.append(company)
                if salary:
                    detail_lines.append(salary)
                if location:
                    detail_lines.append(location)
                if platform:
                    detail_lines.append(f"来源：{platform}")
                detail_str = " · ".join(detail_lines)

                jobs_html += f"""
        <div style="margin-bottom:16px;">
          <table cellpadding="0" cellspacing="0" width="100%">
            <tr>
              <td style="padding-bottom:4px;">
                <span style="font-size:15px;font-weight:700;color:#333;">{title}</span>
                <span style="display:inline-block;margin-left:8px;background:{score_bg};color:{score_color};font-size:11px;font-weight:600;padding:2px 9px;border-radius:10px;">{score}%</span>
              </td>
            </tr>
{f'            <tr><td style="font-size:13px;color:#777;">{detail_str}</td></tr>' if detail_str else ""}
{f'            <tr><td style="font-size:12px;color:#999;margin-top:4px;line-height:1.5;">{reason}</td></tr>' if reason else ""}
{f'            <tr><td style="padding-top:8px;"><a href="{link}" target="_blank" style="color:#9C6B3C;font-size:12px;text-decoration:none;font-weight:600;">查看详情 &rarr;</a></td></tr>' if link else ""}
          </table>
        </div>
        <hr style="border:none;border-top:1px solid #E8DDD0;margin:14px 0;">"""

            login_link_html = (
                f'<div style="margin-top:16px;text-align:center;">'
                f'<a href="{PUBLIC_BASE_URL}/resume_match.html" target="_blank" style="display:inline-block;background:#9C6B3C;color:#fff;text-decoration:none;font-size:13px;font-weight:500;padding:8px 24px;border-radius:6px;">打开就业服务平台查看完整报告</a></div>'
                if PUBLIC_BASE_URL
                else ""
            )

            html_body = f"""
    <div style="max-width:560px;margin:0 auto;font-family:-apple-system,BlinkMacSystemFont,'Segoe UI','PingFang SC','Microsoft YaHei',sans-serif;">
      <div style="background:linear-gradient(135deg,#FBF8F3,#F5E6D3);padding:24px;border-radius:12px 12px 0 0;">
        <h2 style="color:#9C6B3C;margin:0;font-size:20px;letter-spacing:0.5px;">就业服务平台</h2>
        <p style="color:#8B7355;margin:4px 0 0 0;font-size:13px;">岗位智能匹配 · 推荐通知</p>
      </div>
      <div style="padding:24px;border:1px solid #E8DDD0;border-top:none;border-radius:0 0 12px 12px;">
        <p style="margin:0 0 16px;font-size:14px;color:#555;line-height:1.7;">
          {user["username"]}，您好<br>
          您的简历刚刚完成智能匹配分析（{now_str}），发现以下 <b style="color:#9C6B3C;">{job_count}</b> 个与您简历高度匹配的岗位（匹配度 &ge; 60%）
        </p>
        {jobs_html}
        {login_link_html}
        <hr style="border:none;border-top:1px solid #E8DDD0;margin:18px 0 10px;">
        <p style="font-size:11px;color:#AAA;margin:0;line-height:1.6;">此邮件由就业服务平台自动发送 &middot; 请勿直接回复</p>
      </div>
    </div>"""

            msg = MIMEMultipart("alternative")
            msg["Subject"] = "就业服务平台：发现 {} 个高匹配度岗位".format(job_count)
            msg["From"] = email_cfg["from_email"]
            msg["To"] = notify_email
            msg.attach(MIMEText(html_body, "html", "utf-8"))

            # 5. 发送
            import smtplib

            _log(f"Connecting to {email_cfg['smtp_host']}:{email_cfg['smtp_port']}...")
            with smtplib.SMTP_SSL(
                email_cfg["smtp_host"], email_cfg["smtp_port"]
            ) as server:
                server.login(email_cfg["smtp_user"], email_cfg["smtp_pass"])
                server.sendmail(
                    email_cfg["from_email"], [notify_email], msg.as_string()
                )
            _log(
                f"\u2705 Email sent successfully to {notify_email}: {len(high_jobs)} high-prob jobs"
            )

        except Exception as e:
            import traceback

            _log(f"\u274c Failed: {e}")
            traceback.print_exc()
            try:
                with open(log_path, "a", encoding="utf-8") as f:
                    f.write(traceback.format_exc())
            except Exception:
                pass

    t = threading.Thread(target=_do_send, daemon=True)
    t.start()


# ══════════════════════════════════════════════════════
# 辅助函数
# ══════════════════════════════════════════════════════
def check_cache(resume_hash: str, pref_hash: str, job_type: str) -> Optional[dict]:
    """检查匹配缓存"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT result_json, created_at FROM match_cache
        WHERE resume_hash = ? AND preference_hash = ? AND job_type = ?
        ORDER BY created_at DESC LIMIT 1
    """,
        (resume_hash, pref_hash, job_type),
    )
    row = cursor.fetchone()
    conn.close()

    if row:
        result_json, created_at = row["result_json"], row["created_at"]
        try:
            created_dt = datetime.fromisoformat(created_at)
            if (datetime.now() - created_dt).total_seconds() < CACHE_TTL:
                return json.loads(result_json)
        except (ValueError, TypeError):
            pass

    return None


def save_cache(resume_hash: str, pref_hash: str, job_type: str, result: dict):
    """保存匹配结果到缓存"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO match_cache (resume_hash, preference_hash, job_type, result_json)
        VALUES (?, ?, ?, ?)
    """,
        (resume_hash, pref_hash, job_type, json.dumps(result, ensure_ascii=False)),
    )
    conn.commit()
    conn.close()


@app.get("/downloads/{filename}", include_in_schema=False)
async def download_files(filename: str):
    """可下载的工具包（如 Cookie 采集工具）。白名单控制。"""
    _ALLOWED = {"cookie-tool-windows.zip"}
    if filename not in _ALLOWED:
        raise HTTPException(404, "文件不存在")
    path = os.path.join(BASE_DIR, "downloads", filename)
    if not os.path.isfile(path):
        raise HTTPException(404, "工具包未内置到本部署中，请联系管理员获取")
    return FileResponse(path, media_type="application/zip", filename=filename)


@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    """浏览器自动请求 /favicon.ico，用 SVG 图标代替，避免 404 噪声"""
    from fastapi.responses import Response

    svg = (
        "<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' "
        "stroke='#1d1d1f' stroke-width='1.6'>"
        "<rect x='2' y='3' width='8' height='8' rx='1.5'/><rect x='14' y='3' width='8' height='8' rx='1.5'/>"
        "<rect x='2' y='13' width='8' height='8' rx='1.5'/><rect x='14' y='13' width='8' height='8' rx='1.5'/>"
        "</svg>"
    )
    return Response(content=svg, media_type="image/svg+xml")


# ── 前端静态页面（白名单，放在最后避免遮挡 API 路由）─────────────
@app.get("/{filename}", include_in_schema=False)
async def static_pages(filename: str):
    """提供项目根目录下的前端页面与图标（index/resume_match/admin/manifest/图标）

    HTML 页面禁止缓存，保证发新版后用户刷新即得；图标/manifest 可缓存。
    """
    if filename in _STATIC_FILES and os.path.isfile(os.path.join(BASE_DIR, filename)):
        headers = {}
        if filename.endswith(".html"):
            headers["Cache-Control"] = "no-cache"
        return FileResponse(
            os.path.join(BASE_DIR, filename),
            media_type=_STATIC_FILES[filename],
            headers=headers,
        )
    raise HTTPException(404, "页面不存在")


# ── 启动 ─────────────────────────────────────────────
@app.on_event("startup")
async def startup():
    """启动时自动初始化数据库 + 启动定时调度引擎"""
    # 先建岗位表（jobs/gk_jobs/match_cache），再升级用户体系 schema
    import sqlite3 as _sq
    from import_data import create_tables
    from db_schema_v2 import upgrade_schema

    _conn = _sq.connect(DB_PATH)
    create_tables(_conn)
    _conn.close()
    upgrade_schema()
    print("[Startup] 数据库schema检查完成")

    # 启动定时调度（采集 + 订阅推送），单进程内完成，无需独立调度服务
    try:
        import scheduler

        scheduler.start()
    except Exception as e:
        print(f"[Startup] 定时调度启动失败（不影响 Web 服务）: {e}")


if __name__ == "__main__":
    import uvicorn

    port = int(os.environ.get("PORT", "8000"))
    uvicorn.run(app, host="0.0.0.0", port=port)
